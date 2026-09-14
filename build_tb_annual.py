"""Annual TB notification data from UKHSA publications.

1. Local authority x year counts, 2001-2024, at UTLA and LTLA level: UKHSA TB regional reports
   2024, supplementary tables (one workbook per UKHSA region; names only -> matched to ONS codes).
2. UKHSA region x place of birth x year, 2000-2024, with LFS population denominators:
   TB in England 2025 report, Supplementary Table 12.
3. UKHSA region x place of birth x age group x year counts, 2001-2024: regional Table 9.
"""
import re
from pathlib import Path

import openpyxl
import pandas as pd

from build_panel import MERGES, POPULATION_FILES

ROOT = Path(__file__).resolve().parent
RAW = ROOT / "data" / "raw"
SRC = RAW / "tb_sources"
OUT = ROOT / "data" / "processed"

# count table number in each regional workbook (London: boroughs)
TB_SHEETS = {
    "utla": {"East-Midlands": 36, "East-of-England": 38, "London": 39, "North-East": 35, "North-West": 39,
             "South-East": 39, "South-West": 38, "West-Midlands": 36, "Yorkshire-and-the-Humber": 36},
    "ltla": {"East-Midlands": 32, "East-of-England": 34, "London": 35, "North-East": 31, "North-West": 35,
             "South-East": 35, "South-West": 34, "West-Midlands": 32, "Yorkshire-and-the-Humber": 32},
}
REGIONS = list(TB_SHEETS["utla"])


def norm(name):
    s = re.sub(r"\[.*?\]", "", str(name).lower()).replace("&", " and ")
    s = re.sub(r"[^a-z ]", " ", s)
    s = re.sub(r"\b(city of|county of|the|region|statistical)\b", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def sheet_rows(path, table):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    name = next(s for s in wb.sheetnames if s.lower() == f"supplementary_table_{table}".lower())
    return list(wb[name].iter_rows(values_only=True))


def la_annual(level="utla"):
    nomis = pd.read_csv(RAW / POPULATION_FILES[level]).drop_duplicates("GEOGRAPHY_CODE")
    name_to_code = {norm(n): c for n, c in zip(nomis.GEOGRAPHY_NAME, nomis.GEOGRAPHY_CODE)}
    # UKHSA combines these small areas with a neighbour; build_panel.MERGES aligns population, covariates
    # and prescribing on the same neighbour (Hackney; Cornwall) at both levels.
    name_to_code[norm("City of London and Hackney")] = "E09000012"
    name_to_code[norm("Cornwall and Isles of Scilly")] = "E06000052"
    frames = []
    for region, table in TB_SHEETS[level].items():
        path = next((SRC / "regional_2024").glob(f"{region}-TB-report-2024-*.xlsx"))
        rows = sheet_rows(path, table)
        header_i = next(i for i, r in enumerate(rows) if r and str(r[1]).strip() == "2001")
        print(f"{level} {region}: '{rows[0][0]}'")
        df = pd.DataFrame(rows[header_i + 1:], columns=[str(h).strip() for h in rows[header_i]]).dropna(how="all")
        df = df.rename(columns={df.columns[0]: "area_name"})
        long = df.melt(id_vars="area_name", var_name="year", value_name="tb_count")
        long = long[long.year.str.fullmatch(r"\d{4}")]
        long["year"] = long.year.astype(int)
        long["tb_count"] = pd.to_numeric(long.tb_count, errors="coerce")
        long["utla"] = long.area_name.map(norm).map(name_to_code).replace(MERGES[level])
        unmatched = sorted(long.loc[long.utla.isna(), "area_name"].astype(str).unique())
        if unmatched:
            print(f"  unmatched rows (excluded): {unmatched}")
        frames.append(long.dropna(subset=["utla"]).assign(ukhsa_region=region))
    tb = pd.concat(frames).groupby(["utla", "year"], as_index=False).tb_count.sum(min_count=1)
    print(f"{level} x year: {tb.utla.nunique()} areas, {tb.year.min()}-{tb.year.max()}, "
          f"missing counts: {int(tb.tb_count.isna().sum())}")
    print("England totals by year:", tb.groupby("year").tb_count.sum().astype(int).to_dict())
    return tb


def validate_against_fingertips(tb):
    ft = pd.read_csv(RAW / "tb_at502.csv")
    ft = ft[(ft["Indicator ID"] == 91361) & (ft["Area Type"] != "England") & (ft["Time period"] == "2022 - 24")]
    three = tb[tb.year.between(2022, 2024)].groupby("utla").tb_count.sum()
    cmp = ft.set_index("Area Code").Count.to_frame("fingertips_2022_24").join(three.rename("regional_tables_2022_24"))
    cmp = cmp.dropna()
    ratio = cmp.regional_tables_2022_24.sum() / cmp.fingertips_2022_24.sum()
    corr = cmp.corr().iloc[0, 1]
    print(f"Validation vs Fingertips 2022-24 counts ({len(cmp)} UTLAs): total ratio {ratio:.3f}, r = {corr:.4f}, "
          f"max abs difference {int((cmp.regional_tables_2022_24 - cmp.fingertips_2022_24).abs().max())}")
    cmp.to_csv(OUT / "utla_tb_annual_validation.csv")


def region_birthplace():
    rows = sheet_rows(SRC / "TB-in-England-2025-report-supplementary-data-tables-1-incidence-and-epidemiology.xlsx", 12)
    header = [str(h) for h in rows[3]]
    df = pd.DataFrame(rows[4:], columns=header)
    df = df[pd.to_numeric(df.Year, errors="coerce").notna()]
    frames = []
    for col in header:
        m = re.match(r"(UK born|Non-UK born) (.+) number of notifications", col)
        if m:
            birthplace, region = m.groups()
            frames.append(pd.DataFrame({
                "region_name": region, "birthplace": birthplace, "year": df.Year.astype(int),
                "tb_count": pd.to_numeric(df[col], errors="coerce"),
                "population_lfs": pd.to_numeric(df[f"Population: LFS {birthplace} {region}"], errors="coerce"),
            }))
    out = pd.concat(frames)
    print(f"Region x birthplace: regions {sorted(out.region_name.unique())}, {out.year.min()}-{out.year.max()}")
    return out


def region_birthplace_age():
    frames = []
    for region in REGIONS:
        path = next((SRC / "regional_2024").glob(f"{region}-TB-report-2024-*.xlsx"))
        rows = sheet_rows(path, 9)
        header_i = next(i for i, r in enumerate(rows) if r and str(r[0]).strip() == "Year")
        df = pd.DataFrame([r[:4] for r in rows[header_i + 1:]], columns=["year", "age_group", "birthplace", "tb_count"])
        df = df[pd.to_numeric(df.year, errors="coerce").notna()]
        frames.append(df.assign(region_name=region.replace("-", " ")))
    out = pd.concat(frames)
    out["year"] = out.year.astype(int)
    out["tb_count"] = pd.to_numeric(out.tb_count, errors="coerce")
    out["birthplace"] = out.birthplace.str.replace("-", " ").str.replace("Non UK", "Non-UK")
    return out


def main():
    utla = la_annual("utla")
    utla.to_csv(OUT / "utla_tb_annual.csv", index=False)
    validate_against_fingertips(utla)
    la_annual("ltla").to_csv(OUT / "ltla_tb_annual.csv", index=False)
    region_birthplace().to_csv(OUT / "region_tb_birthplace_annual.csv", index=False)
    region_birthplace_age().to_csv(OUT / "region_tb_birthplace_age_annual.csv", index=False)


if __name__ == "__main__":
    main()
